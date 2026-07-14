"""
Script one-shot: migra el histórico ya cargado a mano en
SIBRA-DATA-CORREGIDO/ (Excel) a la Sheet SIBRATECH_INDICES, como semilla
para que los scrapers solo tengan que sumar los datos que faltan desde
ahí en adelante.

Uso:
    cd api
    python migrar_historico.py
"""

from __future__ import annotations

import datetime as dt
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import sheets  # noqa: E402

RAW_DIR = Path(__file__).resolve().parent.parent.parent / "SIBRA-DATA-CORREGIDO" / "RAW"


def _serial_a_fecha(n) -> str | None:
    if not isinstance(n, (int, float)) or n in (0, None):
        return None
    try:
        return (dt.date(1899, 12, 30) + dt.timedelta(days=int(n))).isoformat()
    except (OverflowError, ValueError):
        return None


def _leer_filas(path: Path, hoja: str | None = None):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[hoja or wb.sheetnames[0]]
    return list(ws.iter_rows(values_only=True))


def migrar_cer_uva():
    sid = sheets.ensure_indices_sheet()

    filas_cer = _leer_filas(RAW_DIR / "INDICES" / "CER.xlsx")
    cer = [{"FECHA": _serial_a_fecha(f), "VALOR": v} for f, v, *_ in filas_cer[1:]
           if _serial_a_fecha(f) and v not in (None, "")]
    n = sheets.upsert_series(sid, "CER", ["FECHA", "VALOR"], "FECHA", cer)
    print(f"CER: {n} filas nuevas de {len(cer)} en el histórico")

    filas_uva = _leer_filas(RAW_DIR / "INDICES" / "UV_BCRA.xlsx")
    # la primera fila del archivo original es basura (una sola celda con "1"); el header real es la fila 2
    uva = [{"FECHA": _serial_a_fecha(f), "VALOR": v} for f, v, *_ in filas_uva[2:]
           if _serial_a_fecha(f) and v not in (None, "")]
    n = sheets.upsert_series(sid, "UVA", ["FECHA", "VALOR"], "FECHA", uva)
    print(f"UVA: {n} filas nuevas de {len(uva)} en el histórico")


def migrar_dolar_historico():
    sid = sheets.ensure_indices_sheet()
    filas = _leer_filas(RAW_DIR / "FX" / "USD.xlsx")
    headers = ["FECHA", "OFICIAL_COMPRA", "OFICIAL_VENTA", "BLUE_COMPRA", "BLUE_VENTA",
               "MEP_COMPRA", "MEP_VENTA", "CCL_COMPRA", "CCL_VENTA",
               "MAYORISTA_COMPRA", "MAYORISTA_VENTA", "CRIPTO_COMPRA", "CRIPTO_VENTA",
               "TARJETA_COMPRA", "TARJETA_VENTA"]
    rows = []
    for row in filas[1:]:
        fecha = _serial_a_fecha(row[0])
        of_venta = row[1]
        if not fecha or of_venta in (None, ""):
            continue
        rows.append({"FECHA": fecha, "OFICIAL_VENTA": of_venta})
    n = sheets.upsert_series(sid, "DOLAR", headers, "FECHA", rows)
    print(f"DOLAR (histórico oficial venta): {n} filas nuevas de {len(rows)}")


def migrar_ripte():
    sid = sheets.ensure_indices_sheet()
    filas = _leer_filas(RAW_DIR / "SALARIOS" / "SALARIO.xlsx")
    rows = []
    for row in filas[1:]:
        fecha = _serial_a_fecha(row[0])
        monto = row[1]  # "Monto en $" — coincide con el monto RIPTE oficial (verificado contra la fuente en vivo)
        if not fecha or monto in (None, "", 0):
            continue
        rows.append({"FECHA": fecha, "PERIODO": "", "RIPTE": monto, "VARIACION_MENSUAL": ""})
    n = sheets.upsert_series(sid, "RIPTE", ["FECHA", "PERIODO", "RIPTE", "VARIACION_MENSUAL"], "FECHA", rows)
    print(f"RIPTE: {n} filas nuevas de {len(rows)}")


def migrar_uocra():
    sid = sheets.ensure_indices_sheet()
    filas = _leer_filas(RAW_DIR / "SALARIOS" / "UOCRA.xlsx")
    headers = ["FECHA", "OFICIAL_ESPECIALIZADO", "OFICIAL", "MEDIO_OFICIAL", "AYUDANTE", "SERENO"]
    rows = []
    for row in filas[1:]:
        fecha = _serial_a_fecha(row[0])
        if not fecha:
            continue
        rows.append({
            "FECHA": fecha,
            "OFICIAL_ESPECIALIZADO": row[1],
            "OFICIAL": row[2],
            "MEDIO_OFICIAL": row[3],
            "AYUDANTE": row[4],
            "SERENO": row[5],
        })
    n = sheets.upsert_series(sid, "UOCRA", headers, "FECHA", rows)
    print(f"UOCRA: {n} filas nuevas de {len(rows)}")


def migrar_construccion():
    """CONST_2.xlsx: índice APYMECO base 100 (general/materiales/mano de obra/provisiones) — más
    limpio que CONST.xlsx (que trae CAC/INDEC crudos, con muchos períodos en 0)."""
    sid = sheets.ensure_indices_sheet()
    filas = _leer_filas(RAW_DIR / "CONSTRUCCION" / "CONST_2.xlsx")
    headers = ["FECHA", "INDICE_GENERAL", "MATERIALES", "MANO_DE_OBRA", "PROVISIONES"]
    rows = []
    for row in filas[1:]:
        fecha = _serial_a_fecha(row[0])
        if not fecha or row[1] in (None, ""):
            continue
        rows.append({
            "FECHA": fecha,
            "INDICE_GENERAL": row[1],
            "MATERIALES": row[2],
            "MANO_DE_OBRA": row[3],
            "PROVISIONES": row[4],
        })
    n = sheets.upsert_series(sid, "CONSTRUCCION", headers, "FECHA", rows)
    print(f"CONSTRUCCION: {n} filas nuevas de {len(rows)}")


if __name__ == "__main__":
    print("Migrando histórico de SIBRA-DATA-CORREGIDO a SIBRATECH_INDICES…")
    migrar_cer_uva()
    migrar_dolar_historico()
    migrar_ripte()
    migrar_uocra()
    migrar_construccion()
    print("Listo. Corré server.py y los scrapers se encargan de sumar lo que falta desde acá.")
